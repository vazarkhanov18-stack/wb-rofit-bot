from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import psycopg
from openpyxl import load_workbook
from psycopg.rows import dict_row


DEFAULT_EFFECTIVE_DATE = date(1900, 1, 1)


@dataclass(frozen=True)
class ManagedCostProfile:
    profile_id: int
    sku: str
    name: str
    effective_from: date
    unit_cost: float
    fulfillment_per_unit: float
    packaging_per_unit: float
    warehouse_delivery_per_unit: float
    other_per_unit: float
    created_at: datetime | None = None
    updated_at: datetime | None = None


@dataclass(frozen=True)
class ManagedExpense:
    expense_id: int
    period_start: date
    period_end: date
    category: str
    amount: float
    sku: str
    comment: str
    created_at: datetime | None = None
    updated_at: datetime | None = None


def database_url() -> str:
    return os.getenv("DATABASE_URL", "").strip()


def management_database_enabled() -> bool:
    return bool(database_url())


def _connect():
    url = database_url()
    if not url:
        raise RuntimeError("Не задана переменная DATABASE_URL.")
    return psycopg.connect(url, connect_timeout=12)


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower().replace("ё", "е")
    return re.sub(r"\s+", " ", text)


def normalize_sku(value: Any) -> str:
    return normalize_text(value)


def as_number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    return float(text or 0)


def as_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def init_management_store() -> None:
    with _connect() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS managed_cost_profiles (
                    id BIGSERIAL PRIMARY KEY,
                    sku TEXT NOT NULL,
                    sku_key TEXT NOT NULL,
                    name TEXT NOT NULL DEFAULT '',
                    effective_from DATE NOT NULL,
                    unit_cost DOUBLE PRECISION NOT NULL DEFAULT 0,
                    fulfillment_per_unit DOUBLE PRECISION NOT NULL DEFAULT 0,
                    packaging_per_unit DOUBLE PRECISION NOT NULL DEFAULT 0,
                    warehouse_delivery_per_unit DOUBLE PRECISION NOT NULL DEFAULT 0,
                    other_per_unit DOUBLE PRECISION NOT NULL DEFAULT 0,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    UNIQUE (sku_key, effective_from)
                )
                """
            )
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS managed_external_expenses (
                    id BIGSERIAL PRIMARY KEY,
                    period_start DATE NOT NULL,
                    period_end DATE NOT NULL,
                    category TEXT NOT NULL DEFAULT 'Прочие расходы',
                    amount DOUBLE PRECISION NOT NULL,
                    sku TEXT NOT NULL DEFAULT '',
                    sku_key TEXT NOT NULL DEFAULT '',
                    comment TEXT NOT NULL DEFAULT '',
                    created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    CHECK (period_end >= period_start)
                )
                """
            )
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS management_meta (
                    key TEXT PRIMARY KEY,
                    value TEXT NOT NULL,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            )
            cursor.execute(
                "CREATE INDEX IF NOT EXISTS idx_managed_cost_profiles_sku_date ON managed_cost_profiles(sku_key, effective_from DESC)"
            )
            cursor.execute(
                "CREATE INDEX IF NOT EXISTS idx_managed_expenses_dates ON managed_external_expenses(period_start, period_end)"
            )


def _meta_get(key: str) -> str | None:
    with _connect() as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT value FROM management_meta WHERE key = %s", (key,))
            row = cursor.fetchone()
            return str(row[0]) if row else None


def _meta_set(key: str, value: str) -> None:
    with _connect() as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO management_meta(key, value, updated_at)
                VALUES (%s, %s, NOW())
                ON CONFLICT(key) DO UPDATE SET value = EXCLUDED.value, updated_at = NOW()
                """,
                (key, value),
            )


def management_store_ready() -> bool:
    if not management_database_enabled():
        return False
    try:
        return _meta_get("store_initialized") == "1"
    except Exception:
        return False


def _worksheet_rows(path: Path) -> list[tuple[Any, ...]]:
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        return list(ws.iter_rows(values_only=True))
    finally:
        wb.close()


def _header_map(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    if not rows:
        return 0, {}
    headers = [normalize_text(value) for value in rows[0]]
    return 1, {header: index for index, header in enumerate(headers) if header}


def _value(row: tuple[Any, ...], columns: dict[str, int], *names: str) -> Any:
    for name in names:
        index = columns.get(normalize_text(name))
        if index is not None and index < len(row):
            return row[index]
    return None


def seed_management_from_excel(costs_path: str | Path) -> tuple[int, int]:
    """Однократно переносит старые Excel-справочники в PostgreSQL.

    После успешной попытки PostgreSQL становится главным источником настроек,
    даже если исходные файлы пустые.
    """
    init_management_store()
    if management_store_ready():
        return 0, 0

    costs_file = Path(costs_path)
    data_dir = costs_file.resolve().parent
    history_file = data_dir / "cost_history.xlsx"
    expenses_file = data_dir / "expenses.xlsx"

    imported_costs = 0
    imported_expenses = 0

    # Базовые текущие себестоимости переносим как профиль с 01.01.1900.
    rows = _worksheet_rows(costs_file)
    start, columns = _header_map(rows)
    for row in rows[start:]:
        sku = str(_value(row, columns, "Артикул поставщика", "Артикул") or "").strip()
        if not sku:
            continue
        name = str(_value(row, columns, "Название", "Товар") or sku).strip()
        unit_cost = as_number(_value(row, columns, "себестоимость", "Себестоимость"))
        upsert_cost_profile(
            profile_id=None,
            sku=sku,
            name=name,
            effective_from=DEFAULT_EFFECTIVE_DATE,
            unit_cost=unit_cost,
            fulfillment_per_unit=0,
            packaging_per_unit=0,
            warehouse_delivery_per_unit=0,
            other_per_unit=0,
        )
        imported_costs += 1

    # История цен и расходов на единицу.
    rows = _worksheet_rows(history_file)
    start, columns = _header_map(rows)
    for row in rows[start:]:
        sku = str(_value(row, columns, "Артикул поставщика", "Артикул") or "").strip()
        effective_from = as_date(_value(row, columns, "Действует с", "Дата начала", "Дата"))
        if not sku or effective_from is None:
            continue
        name = str(_value(row, columns, "Название", "Товар") or sku).strip()
        upsert_cost_profile(
            profile_id=None,
            sku=sku,
            name=name,
            effective_from=effective_from,
            unit_cost=as_number(_value(row, columns, "Себестоимость")),
            fulfillment_per_unit=as_number(_value(row, columns, "Фулфилмент на единицу", "Фулфилмент")),
            packaging_per_unit=as_number(_value(row, columns, "Упаковка на единицу", "Упаковка")),
            warehouse_delivery_per_unit=as_number(_value(row, columns, "Доставка до склада на единицу", "Доставка до склада")),
            other_per_unit=as_number(_value(row, columns, "Прочие расходы на единицу", "Прочие на единицу")),
        )
        imported_costs += 1

    # Общие внешние расходы.
    rows = _worksheet_rows(expenses_file)
    start, columns = _header_map(rows)
    for row in rows[start:]:
        period_start = as_date(_value(row, columns, "Дата начала", "Дата расхода", "Дата"))
        if period_start is None:
            continue
        period_end = as_date(_value(row, columns, "Дата окончания", "Дата конца")) or period_start
        amount = as_number(_value(row, columns, "Сумма", "Расход"))
        if abs(amount) < 1e-9:
            continue
        upsert_expense(
            expense_id=None,
            period_start=period_start,
            period_end=period_end,
            category=str(_value(row, columns, "Категория", "Тип расхода") or "Прочие расходы").strip(),
            amount=amount,
            sku=str(_value(row, columns, "Артикул поставщика", "Артикул") or "").strip(),
            comment=str(_value(row, columns, "Комментарий", "Описание") or "").strip(),
        )
        imported_expenses += 1

    _meta_set("store_initialized", "1")
    _meta_set("seeded_from_excel_at", datetime.utcnow().isoformat(timespec="seconds"))
    return imported_costs, imported_expenses


def list_cost_profiles(limit: int = 1000) -> list[ManagedCostProfile]:
    limit = max(1, min(int(limit), 5000))
    with _connect() as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute(
                """
                SELECT id AS profile_id, sku, name, effective_from, unit_cost,
                       fulfillment_per_unit, packaging_per_unit,
                       warehouse_delivery_per_unit, other_per_unit,
                       created_at, updated_at
                FROM managed_cost_profiles
                ORDER BY sku_key ASC, effective_from DESC, id DESC
                LIMIT %s
                """,
                (limit,),
            )
            return [ManagedCostProfile(**dict(row)) for row in cursor.fetchall()]


def get_cost_profile(profile_id: int) -> ManagedCostProfile | None:
    with _connect() as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute(
                """
                SELECT id AS profile_id, sku, name, effective_from, unit_cost,
                       fulfillment_per_unit, packaging_per_unit,
                       warehouse_delivery_per_unit, other_per_unit,
                       created_at, updated_at
                FROM managed_cost_profiles WHERE id = %s
                """,
                (int(profile_id),),
            )
            row = cursor.fetchone()
            return ManagedCostProfile(**dict(row)) if row else None


def upsert_cost_profile(
    *,
    profile_id: int | None,
    sku: str,
    name: str,
    effective_from: date,
    unit_cost: float,
    fulfillment_per_unit: float = 0,
    packaging_per_unit: float = 0,
    warehouse_delivery_per_unit: float = 0,
    other_per_unit: float = 0,
) -> int:
    sku = str(sku or "").strip()
    if not sku:
        raise ValueError("Укажи артикул поставщика.")
    sku_key = normalize_sku(sku)
    name = str(name or sku).strip() or sku
    values = tuple(float(value) for value in (
        unit_cost,
        fulfillment_per_unit,
        packaging_per_unit,
        warehouse_delivery_per_unit,
        other_per_unit,
    ))
    if any(value < 0 for value in values):
        raise ValueError("Себестоимость и расходы на единицу не могут быть отрицательными.")

    with _connect() as connection:
        with connection.cursor() as cursor:
            if profile_id:
                cursor.execute(
                    """
                    UPDATE managed_cost_profiles SET
                        sku=%s, sku_key=%s, name=%s, effective_from=%s,
                        unit_cost=%s, fulfillment_per_unit=%s,
                        packaging_per_unit=%s, warehouse_delivery_per_unit=%s,
                        other_per_unit=%s, updated_at=NOW()
                    WHERE id=%s
                    RETURNING id
                    """,
                    (sku, sku_key, name, effective_from, *values, int(profile_id)),
                )
                row = cursor.fetchone()
                if not row:
                    raise ValueError("Запись себестоимости не найдена.")
                return int(row[0])

            cursor.execute(
                """
                INSERT INTO managed_cost_profiles (
                    sku, sku_key, name, effective_from, unit_cost,
                    fulfillment_per_unit, packaging_per_unit,
                    warehouse_delivery_per_unit, other_per_unit,
                    created_at, updated_at
                ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,NOW(),NOW())
                ON CONFLICT (sku_key, effective_from) DO UPDATE SET
                    sku=EXCLUDED.sku, name=EXCLUDED.name,
                    unit_cost=EXCLUDED.unit_cost,
                    fulfillment_per_unit=EXCLUDED.fulfillment_per_unit,
                    packaging_per_unit=EXCLUDED.packaging_per_unit,
                    warehouse_delivery_per_unit=EXCLUDED.warehouse_delivery_per_unit,
                    other_per_unit=EXCLUDED.other_per_unit,
                    updated_at=NOW()
                RETURNING id
                """,
                (sku, sku_key, name, effective_from, *values),
            )
            row = cursor.fetchone()
            return int(row[0])


def delete_cost_profile(profile_id: int) -> bool:
    with _connect() as connection:
        with connection.cursor() as cursor:
            cursor.execute("DELETE FROM managed_cost_profiles WHERE id=%s", (int(profile_id),))
            return cursor.rowcount > 0


def list_expenses(limit: int = 1000) -> list[ManagedExpense]:
    limit = max(1, min(int(limit), 5000))
    with _connect() as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute(
                """
                SELECT id AS expense_id, period_start, period_end, category,
                       amount, sku, comment, created_at, updated_at
                FROM managed_external_expenses
                ORDER BY period_start DESC, id DESC
                LIMIT %s
                """,
                (limit,),
            )
            return [ManagedExpense(**dict(row)) for row in cursor.fetchall()]


def get_expense(expense_id: int) -> ManagedExpense | None:
    with _connect() as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute(
                """
                SELECT id AS expense_id, period_start, period_end, category,
                       amount, sku, comment, created_at, updated_at
                FROM managed_external_expenses WHERE id=%s
                """,
                (int(expense_id),),
            )
            row = cursor.fetchone()
            return ManagedExpense(**dict(row)) if row else None


def upsert_expense(
    *,
    expense_id: int | None,
    period_start: date,
    period_end: date,
    category: str,
    amount: float,
    sku: str = "",
    comment: str = "",
) -> int:
    if period_end < period_start:
        raise ValueError("Дата окончания не может быть раньше даты начала.")
    amount = float(amount)
    if abs(amount) < 1e-9:
        raise ValueError("Сумма расхода не может быть нулевой.")
    category = str(category or "Прочие расходы").strip() or "Прочие расходы"
    sku = str(sku or "").strip()
    sku_key = normalize_sku(sku)
    comment = str(comment or "").strip()

    with _connect() as connection:
        with connection.cursor() as cursor:
            if expense_id:
                cursor.execute(
                    """
                    UPDATE managed_external_expenses SET
                        period_start=%s, period_end=%s, category=%s, amount=%s,
                        sku=%s, sku_key=%s, comment=%s, updated_at=NOW()
                    WHERE id=%s RETURNING id
                    """,
                    (period_start, period_end, category, amount, sku, sku_key, comment, int(expense_id)),
                )
                row = cursor.fetchone()
                if not row:
                    raise ValueError("Расход не найден.")
                return int(row[0])

            cursor.execute(
                """
                INSERT INTO managed_external_expenses (
                    period_start, period_end, category, amount,
                    sku, sku_key, comment, created_at, updated_at
                ) VALUES (%s,%s,%s,%s,%s,%s,%s,NOW(),NOW())
                RETURNING id
                """,
                (period_start, period_end, category, amount, sku, sku_key, comment),
            )
            row = cursor.fetchone()
            return int(row[0])


def delete_expense(expense_id: int) -> bool:
    with _connect() as connection:
        with connection.cursor() as cursor:
            cursor.execute("DELETE FROM managed_external_expenses WHERE id=%s", (int(expense_id),))
            return cursor.rowcount > 0


def management_counts() -> tuple[int, int]:
    with _connect() as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) FROM managed_cost_profiles")
            cost_count = int(cursor.fetchone()[0])
            cursor.execute("SELECT COUNT(*) FROM managed_external_expenses")
            expense_count = int(cursor.fetchone()[0])
            return cost_count, expense_count
