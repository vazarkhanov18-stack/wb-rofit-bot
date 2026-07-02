from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable
import math
import re

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


TAX_RATE = 0.06
UNALLOCATED_SKU = "НЕРАСПРЕДЕЛЕНО"
DEFAULT_EFFECTIVE_DATE = date(1900, 1, 1)


@dataclass(frozen=True)
class CostItem:
    sku: str
    name: str
    unit_cost: float


@dataclass(frozen=True)
class CostProfile:
    sku: str
    name: str
    effective_from: date
    unit_cost: float = 0.0
    fulfillment_per_unit: float = 0.0
    packaging_per_unit: float = 0.0
    warehouse_delivery_per_unit: float = 0.0
    other_per_unit: float = 0.0

    @property
    def unit_external_expense(self) -> float:
        return (
            self.fulfillment_per_unit
            + self.packaging_per_unit
            + self.warehouse_delivery_per_unit
            + self.other_per_unit
        )


@dataclass(frozen=True)
class ExternalExpense:
    period_start: date
    period_end: date
    category: str
    amount: float
    sku: str = ""
    comment: str = ""


@dataclass
class SkuResult:
    sku: str
    name: str = ""
    unit_cost: float = 0.0
    nm_id: int | None = None
    sold_units: float = 0.0
    revenue: float = 0.0
    payout_for_goods: float = 0.0
    logistics: float = 0.0
    transport: float = 0.0
    handling: float = 0.0
    storage: float = 0.0
    other_withholdings: float = 0.0
    fines: float = 0.0
    advertising: float = 0.0
    cogs_amount: float = 0.0
    unit_external_expenses: float = 0.0
    general_external_expenses: float = 0.0

    @property
    def calculated_payout(self) -> float:
        return (
            self.payout_for_goods
            - self.logistics
            - self.transport
            - self.handling
            - self.storage
            - self.other_withholdings
            - self.fines
        )

    @property
    def cogs(self) -> float:
        return self.cogs_amount

    @property
    def external_expenses(self) -> float:
        return self.unit_external_expenses + self.general_external_expenses

    @property
    def tax(self) -> float:
        return self.revenue * TAX_RATE

    @property
    def operating_profit_before_ads_and_tax(self) -> float:
        return self.calculated_payout - self.cogs - self.external_expenses

    @property
    def profit_before_ads(self) -> float:
        # Сохраняем прежнее поле для совместимости с БД: прибыль после налога, но до рекламы.
        return self.operating_profit_before_ads_and_tax - self.tax

    @property
    def profit_before_tax(self) -> float:
        return self.operating_profit_before_ads_and_tax - self.advertising

    @property
    def profit(self) -> float:
        return self.profit_before_tax - self.tax

    @property
    def margin(self) -> float:
        return self.profit / self.revenue if self.revenue else 0.0

    @property
    def drr(self) -> float:
        return self.advertising / self.revenue if self.revenue else 0.0

    @property
    def profit_per_unit(self) -> float:
        return self.profit / self.sold_units if self.sold_units else 0.0

    @property
    def has_activity(self) -> bool:
        return any(
            abs(v) > 1e-9
            for v in (
                self.sold_units,
                self.revenue,
                self.payout_for_goods,
                self.logistics,
                self.transport,
                self.handling,
                self.storage,
                self.other_withholdings,
                self.fines,
                self.advertising,
                self.cogs_amount,
                self.unit_external_expenses,
                self.general_external_expenses,
            )
        )


@dataclass
class ReportResult:
    period_start: date | None
    period_end: date | None
    items: list[SkuResult]
    missing_cost_skus: list[str] = field(default_factory=list)
    unmatched_ad_nm_ids: list[int] = field(default_factory=list)
    advertising_warning: str = ""
    advertising_campaign_count: int = 0
    external_expense_warning: str = ""

    @property
    def active_items(self) -> list[SkuResult]:
        return [item for item in self.items if item.has_activity]

    def total(self) -> SkuResult:
        total = SkuResult(sku="ИТОГО", name="Итого")
        for item in self.active_items:
            total.sold_units += item.sold_units
            total.revenue += item.revenue
            total.payout_for_goods += item.payout_for_goods
            total.logistics += item.logistics
            total.transport += item.transport
            total.handling += item.handling
            total.storage += item.storage
            total.other_withholdings += item.other_withholdings
            total.fines += item.fines
            total.advertising += item.advertising
            total.cogs_amount += item.cogs
            total.unit_external_expenses += item.unit_external_expenses
            total.general_external_expenses += item.general_external_expenses
        return total

    @property
    def total_cogs(self) -> float:
        return sum(item.cogs for item in self.active_items)

    @property
    def total_unit_external_expenses(self) -> float:
        return sum(item.unit_external_expenses for item in self.active_items)

    @property
    def total_general_external_expenses(self) -> float:
        return sum(item.general_external_expenses for item in self.active_items)

    @property
    def total_external_expenses(self) -> float:
        return self.total_unit_external_expenses + self.total_general_external_expenses

    @property
    def total_operating_profit_before_ads_and_tax(self) -> float:
        return sum(item.operating_profit_before_ads_and_tax for item in self.active_items)

    @property
    def total_profit_before_ads(self) -> float:
        return sum(item.profit_before_ads for item in self.active_items)

    @property
    def total_profit_before_tax(self) -> float:
        return sum(item.profit_before_tax for item in self.active_items)

    @property
    def total_advertising(self) -> float:
        return sum(item.advertising for item in self.active_items)

    @property
    def total_profit(self) -> float:
        return sum(item.profit for item in self.active_items)

    @property
    def total_revenue(self) -> float:
        return sum(item.revenue for item in self.active_items)

    @property
    def total_units(self) -> float:
        return sum(item.sold_units for item in self.active_items)

    @property
    def total_payout(self) -> float:
        return sum(item.calculated_payout for item in self.active_items)

    @property
    def total_tax(self) -> float:
        return sum(item.tax for item in self.active_items)

    @property
    def margin(self) -> float:
        return self.total_profit / self.total_revenue if self.total_revenue else 0.0

    @property
    def drr(self) -> float:
        return self.total_advertising / self.total_revenue if self.total_revenue else 0.0

    @property
    def profit_per_unit(self) -> float:
        return self.total_profit / self.total_units if self.total_units else 0.0


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "sku": ("Артикул поставщика",),
    "name": ("Название",),
    "barcode": ("Баркод",),
    "reason": ("Обоснование для оплаты",),
    "date": ("Дата продажи", "Дата заказа покупателем"),
    "qty": ("Кол-во",),
    "revenue": ("Вайлдберриз реализовал Товар (Пр)",),
    "payout": ("К перечислению Продавцу за реализованный Товар",),
    "logistics": ("Услуги по доставке товара покупателю",),
    "transport": ("Возмещение издержек по перевозке/по складским операциям с товаром",),
    "handling": ("Операции на приемке",),
    "storage": ("Хранение",),
    "withholdings": ("Удержания",),
    "fines": ("Общая сумма штрафов",),
}


COST_HISTORY_ALIASES: dict[str, tuple[str, ...]] = {
    "sku": ("Артикул поставщика", "Артикул"),
    "name": ("Название", "Товар"),
    "effective_from": ("Действует с", "Дата начала", "Дата"),
    "unit_cost": ("Себестоимость", "себестоимость"),
    "fulfillment": ("Фулфилмент на единицу", "Фулфилмент"),
    "packaging": ("Упаковка на единицу", "Упаковка"),
    "warehouse_delivery": ("Доставка до склада на единицу", "Доставка до склада"),
    "other": ("Прочие расходы на единицу", "Прочие на единицу"),
}


EXPENSE_ALIASES: dict[str, tuple[str, ...]] = {
    "start": ("Дата начала", "Дата расхода", "Дата"),
    "end": ("Дата окончания", "Дата конца"),
    "category": ("Категория", "Тип расхода"),
    "amount": ("Сумма", "Расход"),
    "sku": ("Артикул поставщика", "Артикул"),
    "comment": ("Комментарий", "Описание"),
}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower().replace("ё", "е")
    return re.sub(r"\s+", " ", text)


def normalize_header(value: Any) -> str:
    return normalize_text(value).replace("\n", " ")


def normalize_sku(value: Any) -> str:
    return normalize_text(value)


def valid_sku(value: Any) -> bool:
    sku = normalize_sku(value)
    return sku not in {"", "0", "90", "none", "nan", "null"}


def as_number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return 0.0
        return float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    if text.lower() in {"", "none", "nan", "90"}:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (ValueError, TypeError, OverflowError):
            pass
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def find_header_row(rows: Iterable[tuple[Any, ...]], required_header: str, max_rows: int = 15) -> tuple[int, list[Any]]:
    target = normalize_header(required_header)
    for index, row in enumerate(rows, start=1):
        normalized = [normalize_header(cell) for cell in row]
        if target in normalized:
            return index, list(row)
        if index >= max_rows:
            break
    raise ValueError(f"Не найдена строка заголовков с колонкой «{required_header}».")


def _resolve_alias_columns(headers: list[Any], aliases: dict[str, tuple[str, ...]]) -> dict[str, int]:
    normalized = [normalize_header(h) for h in headers]
    result: dict[str, int] = {}
    for key, variants in aliases.items():
        for variant in variants:
            target = normalize_header(variant)
            if target in normalized:
                result[key] = normalized.index(target)
                break
    return result


def resolve_columns(headers: list[Any]) -> dict[str, int]:
    result = _resolve_alias_columns(headers, HEADER_ALIASES)
    required = ("sku", "reason", "qty", "revenue", "payout")
    missing = [key for key in required if key not in result]
    if missing:
        raise ValueError("В отчёте не хватает обязательных колонок: " + ", ".join(missing))
    return result


def load_costs(costs_path: str | Path) -> dict[str, CostItem]:
    """Читает прежний costs.xlsx. Сохранено для обратной совместимости."""
    path = Path(costs_path)
    if not path.exists():
        raise FileNotFoundError(f"Не найден файл себестоимости: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    _, headers = find_header_row(iter(all_rows), "Артикул поставщика")
    columns = _resolve_alias_columns(headers, COST_HISTORY_ALIASES)
    if "sku" not in columns or "unit_cost" not in columns:
        wb.close()
        raise ValueError("В файле себестоимости нужны колонки «Артикул поставщика» и «себестоимость».")

    header_index, _ = find_header_row(iter(all_rows), "Артикул поставщика")
    costs: dict[str, CostItem] = {}
    for row in all_rows[header_index:]:
        sku_col = columns["sku"]
        if sku_col >= len(row) or not valid_sku(row[sku_col]):
            continue
        sku_raw = str(row[sku_col]).strip()
        key = normalize_sku(sku_raw)
        name_col = columns.get("name")
        name = (
            str(row[name_col] or sku_raw).strip()
            if name_col is not None and name_col < len(row)
            else sku_raw
        )
        cost_col = columns["unit_cost"]
        cost = as_number(row[cost_col]) if cost_col < len(row) else 0.0
        costs[key] = CostItem(sku=sku_raw, name=name, unit_cost=cost)
    wb.close()
    return costs


def _cost_history_path(costs_path: str | Path) -> Path:
    return Path(costs_path).resolve().parent / "cost_history.xlsx"


def _expenses_path(costs_path: str | Path) -> Path:
    return Path(costs_path).resolve().parent / "expenses.xlsx"


def load_cost_profiles(costs_path: str | Path) -> dict[str, list[CostProfile]]:
    """Объединяет базовый costs.xlsx с необязательной историей cost_history.xlsx."""
    baseline = load_costs(costs_path)
    profiles: dict[str, list[CostProfile]] = {
        key: [
            CostProfile(
                sku=item.sku,
                name=item.name,
                effective_from=DEFAULT_EFFECTIVE_DATE,
                unit_cost=item.unit_cost,
            )
        ]
        for key, item in baseline.items()
    }

    path = _cost_history_path(costs_path)
    if not path.exists():
        return profiles

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        wb.close()
        return profiles

    try:
        header_index, headers = find_header_row(iter(all_rows), "Артикул поставщика")
    except ValueError:
        wb.close()
        return profiles
    columns = _resolve_alias_columns(headers, COST_HISTORY_ALIASES)
    if "sku" not in columns or "effective_from" not in columns or "unit_cost" not in columns:
        wb.close()
        return profiles

    for row in all_rows[header_index:]:
        sku_col = columns["sku"]
        if sku_col >= len(row) or not valid_sku(row[sku_col]):
            continue
        effective_col = columns["effective_from"]
        effective = as_date(row[effective_col]) if effective_col < len(row) else None
        if effective is None:
            continue
        sku_raw = str(row[sku_col]).strip()
        key = normalize_sku(sku_raw)
        name_col = columns.get("name")
        name = (
            str(row[name_col] or sku_raw).strip()
            if name_col is not None and name_col < len(row)
            else sku_raw
        )

        def number_for(column_key: str) -> float:
            idx = columns.get(column_key)
            return as_number(row[idx]) if idx is not None and idx < len(row) else 0.0

        profiles.setdefault(key, []).append(
            CostProfile(
                sku=sku_raw,
                name=name,
                effective_from=effective,
                unit_cost=number_for("unit_cost"),
                fulfillment_per_unit=number_for("fulfillment"),
                packaging_per_unit=number_for("packaging"),
                warehouse_delivery_per_unit=number_for("warehouse_delivery"),
                other_per_unit=number_for("other"),
            )
        )
    wb.close()

    for key in profiles:
        # Если на одну дату несколько строк, последняя строка файла имеет приоритет.
        dedup: dict[date, CostProfile] = {}
        for profile in profiles[key]:
            dedup[profile.effective_from] = profile
        profiles[key] = sorted(dedup.values(), key=lambda item: item.effective_from)
    return profiles


def resolve_cost_profile(
    profiles: dict[str, list[CostProfile]],
    sku_key: str,
    transaction_date: date | None,
) -> CostProfile | None:
    history = profiles.get(normalize_sku(sku_key))
    if not history:
        return None
    target = transaction_date or date.max
    eligible = [profile for profile in history if profile.effective_from <= target]
    return eligible[-1] if eligible else history[0]


def load_external_expenses(costs_path: str | Path) -> list[ExternalExpense]:
    path = _expenses_path(costs_path)
    if not path.exists():
        return []

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        wb.close()
        return []

    try:
        header_index, headers = find_header_row(iter(all_rows), "Дата начала")
    except ValueError:
        wb.close()
        return []
    columns = _resolve_alias_columns(headers, EXPENSE_ALIASES)
    if "start" not in columns or "amount" not in columns:
        wb.close()
        return []

    expenses: list[ExternalExpense] = []
    for row in all_rows[header_index:]:
        start_col = columns["start"]
        start = as_date(row[start_col]) if start_col < len(row) else None
        if start is None:
            continue
        end_col = columns.get("end")
        end = as_date(row[end_col]) if end_col is not None and end_col < len(row) else None
        end = end or start
        if end < start:
            start, end = end, start
        amount_col = columns["amount"]
        amount = as_number(row[amount_col]) if amount_col < len(row) else 0.0
        if abs(amount) < 1e-9:
            continue
        category_col = columns.get("category")
        category = (
            str(row[category_col] or "Прочие расходы").strip()
            if category_col is not None and category_col < len(row)
            else "Прочие расходы"
        )
        sku_col = columns.get("sku")
        sku = (
            str(row[sku_col] or "").strip()
            if sku_col is not None and sku_col < len(row)
            else ""
        )
        comment_col = columns.get("comment")
        comment = (
            str(row[comment_col] or "").strip()
            if comment_col is not None and comment_col < len(row)
            else ""
        )
        expenses.append(
            ExternalExpense(
                period_start=start,
                period_end=end,
                category=category,
                amount=amount,
                sku=sku,
                comment=comment,
            )
        )
    wb.close()
    return expenses


def _overlap_share(expense: ExternalExpense, report_start: date, report_end: date) -> float:
    overlap_start = max(expense.period_start, report_start)
    overlap_end = min(expense.period_end, report_end)
    if overlap_end < overlap_start:
        return 0.0
    expense_days = (expense.period_end - expense.period_start).days + 1
    overlap_days = (overlap_end - overlap_start).days + 1
    return expense.amount * overlap_days / max(1, expense_days)


def apply_external_expenses(result: ReportResult, costs_path: str | Path) -> ReportResult:
    if result.period_start is None or result.period_end is None:
        return result
    expenses = load_external_expenses(costs_path)
    if not expenses:
        return result

    items_by_sku = {normalize_sku(item.sku): item for item in result.items if item.sku}
    general_amount = 0.0
    warnings: list[str] = []

    for expense in expenses:
        amount = _overlap_share(expense, result.period_start, result.period_end)
        if abs(amount) < 1e-9:
            continue
        if valid_sku(expense.sku):
            key = normalize_sku(expense.sku)
            item = items_by_sku.get(key)
            if item is None:
                item = SkuResult(
                    sku=expense.sku,
                    name=f"{expense.category}: {expense.sku}",
                )
                result.items.append(item)
                items_by_sku[key] = item
            item.general_external_expenses += amount
        else:
            general_amount += amount

    if abs(general_amount) > 1e-9:
        allocatable = [item for item in result.items if item.revenue > 0]
        total_revenue = sum(item.revenue for item in allocatable)
        if total_revenue > 0:
            allocated = 0.0
            for item in allocatable[:-1]:
                share = general_amount * item.revenue / total_revenue
                item.general_external_expenses += share
                allocated += share
            allocatable[-1].general_external_expenses += general_amount - allocated
        else:
            key = normalize_sku("ВНЕШНИЕ-НЕРАСПРЕДЕЛЕНО")
            item = items_by_sku.get(key)
            if item is None:
                item = SkuResult(
                    sku="ВНЕШНИЕ-НЕРАСПРЕДЕЛЕНО",
                    name="Общие внешние расходы без продаж",
                )
                result.items.append(item)
                items_by_sku[key] = item
            item.general_external_expenses += general_amount
            warnings.append("Общие внешние расходы не удалось распределить по товарам: за период нет положительного дохода.")

    result.external_expense_warning = " ".join(warnings)
    result.items = [item for item in result.items if item.has_activity]
    result.items.sort(key=lambda item: (item.revenue, item.profit), reverse=True)
    return result


def analyze_report(report_path: str | Path, costs_path: str | Path) -> ReportResult:
    profiles = load_cost_profiles(costs_path)
    wb = load_workbook(report_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    all_rows = list(ws.iter_rows(values_only=True))
    header_index, headers = find_header_row(iter(all_rows), "Обоснование для оплаты")
    columns = resolve_columns(headers)
    data_rows = all_rows[header_index:]

    barcode_to_sku: dict[str, str] = {}
    report_names: dict[str, str] = {}
    for row in data_rows:
        sku_value = row[columns["sku"]] if columns["sku"] < len(row) else None
        barcode_index = columns.get("barcode", -1)
        barcode_value = row[barcode_index] if barcode_index >= 0 and barcode_index < len(row) else None
        if valid_sku(sku_value):
            sku_key = normalize_sku(sku_value)
            if barcode_value not in (None, ""):
                barcode_to_sku[str(barcode_value).strip()] = sku_key
            if "name" in columns and columns["name"] < len(row):
                name = str(row[columns["name"]] or "").strip()
                if name and name != "90":
                    report_names[sku_key] = name

    results: dict[str, SkuResult] = {}
    for key, history in profiles.items():
        latest = history[-1]
        results[key] = SkuResult(sku=latest.sku, name=latest.name, unit_cost=latest.unit_cost)

    dates: list[date] = []
    missing_keys: set[str] = set()

    for row in data_rows:
        if not any(cell not in (None, "") for cell in row):
            continue

        reason = normalize_text(row[columns["reason"]] if columns["reason"] < len(row) else None)
        is_sale = reason == "продажа" or reason.startswith("продажа ")
        is_return = "возврат" in reason and reason != "возмещение издержек по перевозке/по складским операциям с товаром"

        sku_value = row[columns["sku"]] if columns["sku"] < len(row) else None
        sku_key = normalize_sku(sku_value) if valid_sku(sku_value) else ""
        if not sku_key and (is_sale or is_return) and "barcode" in columns:
            barcode_value = row[columns["barcode"]] if columns["barcode"] < len(row) else None
            sku_key = barcode_to_sku.get(str(barcode_value).strip(), "")
        if not sku_key:
            sku_key = normalize_sku(UNALLOCATED_SKU)

        if sku_key not in results:
            display_sku = str(sku_value).strip() if valid_sku(sku_value) else UNALLOCATED_SKU
            display_name = report_names.get(
                sku_key,
                "Расходы без артикула в отчёте WB" if display_sku == UNALLOCATED_SKU else display_sku,
            )
            results[sku_key] = SkuResult(sku=display_sku, name=display_name)

        item = results[sku_key]
        if not item.name:
            item.name = report_names.get(sku_key, item.sku)

        transaction_date: date | None = None
        if "date" in columns and columns["date"] < len(row):
            transaction_date = as_date(row[columns["date"]])
            if transaction_date:
                dates.append(transaction_date)

        if is_sale or is_return:
            qty = as_number(row[columns["qty"]])
            revenue = as_number(row[columns["revenue"]])
            payout = as_number(row[columns["payout"]])
            if is_return:
                qty = -abs(qty) if qty >= 0 else qty
                revenue = -abs(revenue) if revenue >= 0 else revenue
                payout = -abs(payout) if payout >= 0 else payout

            item.sold_units += qty
            item.revenue += revenue
            item.payout_for_goods += payout
            profile = resolve_cost_profile(profiles, sku_key, transaction_date)
            if profile is None:
                missing_keys.add(item.sku)
            else:
                item.unit_cost = profile.unit_cost
                if not item.name:
                    item.name = profile.name
                item.cogs_amount += qty * profile.unit_cost
                item.unit_external_expenses += qty * profile.unit_external_expense

        optional_fields = {
            "logistics": "logistics",
            "transport": "transport",
            "handling": "handling",
            "storage": "storage",
            "withholdings": "other_withholdings",
            "fines": "fines",
        }
        for column_key, attribute in optional_fields.items():
            if column_key in columns and columns[column_key] < len(row):
                setattr(item, attribute, getattr(item, attribute) + as_number(row[columns[column_key]]))

    wb.close()

    result = ReportResult(
        period_start=min(dates) if dates else None,
        period_end=max(dates) if dates else None,
        items=[item for item in results.values() if item.has_activity],
        missing_cost_skus=sorted(missing_keys),
    )
    apply_external_expenses(result, costs_path)
    result.items.sort(key=lambda item: (item.revenue, item.profit), reverse=True)
    return result


def format_money(value: float) -> str:
    sign = "−" if value < 0 else ""
    number = f"{abs(value):,.2f}".replace(",", " ")
    return f"{sign}{number} ₽"


def format_units(value: float) -> str:
    if abs(value - round(value)) < 1e-9:
        return str(int(round(value)))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def format_period(start: date | None, end: date | None) -> str:
    if not start and not end:
        return "период не определён"
    if start == end or end is None:
        return start.strftime("%d.%m.%Y") if start else end.strftime("%d.%m.%Y")
    return f"{start.strftime('%d.%m.%Y')}–{end.strftime('%d.%m.%Y')}"


def build_messages(result: ReportResult) -> list[str]:
    total_profit = result.total_profit
    total_revenue = result.total_revenue

    summary = [
        f"📊 Отчёт WB за {format_period(result.period_start, result.period_end)}",
        "",
        f"Продано: {format_units(result.total_units)} шт.",
        f"Доход покупателей: {format_money(total_revenue)}",
        f"Расчётная выплата: {format_money(result.total_payout)}",
        f"Себестоимость: {format_money(result.total_cogs)}",
        f"Внешние расходы: {format_money(result.total_external_expenses)}",
        f"Прибыль до рекламы и налога: {format_money(result.total_operating_profit_before_ads_and_tax)}",
        f"Реклама WB: {format_money(result.total_advertising)}",
        f"ДРР: {result.drr * 100:.1f}%",
        f"Прибыль до налога: {format_money(result.total_profit_before_tax)}",
        f"УСН 6%: {format_money(result.total_tax)}",
        "",
        f"💰 Чистая прибыль: {format_money(total_profit)}",
        f"Маржинальность: {result.margin * 100:.1f}%",
        f"Прибыль на единицу: {format_money(result.profit_per_unit)}",
    ]

    if result.advertising_warning:
        summary.extend(["", "⚠️ Реклама не была учтена полностью:", result.advertising_warning])

    if result.external_expense_warning:
        summary.extend(["", "⚠️ Внешние расходы:", result.external_expense_warning])

    if result.missing_cost_skus:
        summary.extend(
            [
                "",
                "⚠️ Не найдена себестоимость: " + ", ".join(result.missing_cost_skus),
                "Итоговая прибыль по этим товарам может быть завышена.",
            ]
        )

    if result.unmatched_ad_nm_ids:
        shown = ", ".join(str(value) for value in result.unmatched_ad_nm_ids[:10])
        if len(result.unmatched_ad_nm_ids) > 10:
            shown += "…"
        summary.extend(
            [
                "",
                "⚠️ Реклама найдена, но не связана с артикулом продавца для WB-артикулов: " + shown,
                "Общая прибыль учтена верно, но расходы по товарам могут отображаться отдельными строками.",
            ]
        )

    detail = ["📦 По товарам:"]
    for item in result.active_items:
        label = item.name or item.sku
        if len(label) > 55:
            label = label[:52] + "…"
        detail.extend(
            [
                "",
                label,
                f"Артикул: {item.sku}",
                f"Продано: {format_units(item.sold_units)} | Доход: {format_money(item.revenue)}",
                f"Внешние: {format_money(item.external_expenses)} | Реклама: {format_money(item.advertising)}",
                f"До налога: {format_money(item.profit_before_tax)} | УСН: {format_money(item.tax)}",
                f"Чистая прибыль: {format_money(item.profit)} | Маржа: {item.margin * 100:.1f}%",
            ]
        )

    notes = [
        "",
        "Расчёт предварительный: учтены финансовый отчёт WB, история себестоимости, расходы на единицу, внешние расходы, УСН 6% и WB Продвижение. Точность зависит от заполнения файлов cost_history.xlsx и expenses.xlsx.",
    ]

    return ["\n".join(summary), "\n".join(detail + notes)]
